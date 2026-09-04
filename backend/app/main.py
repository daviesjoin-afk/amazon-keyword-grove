"""FastAPI entrypoint for the local Amazon keyword library MVP."""

from __future__ import annotations

import csv
import io
import math
import os
import threading
from typing import Any, Callable

from fastapi import Body, FastAPI, File, Form, HTTPException, Query, UploadFile
from fastapi.middleware.cors import CORSMiddleware
from fastapi.responses import StreamingResponse

from .ai_service import (
    ai_json_response as _ai_json_response,
    is_short_generic_query as _is_short_generic_query,
    public_ai_config as _public_ai_config,
    require_ai_config as _ai_config_or_error,
    run_semantic_batches as _run_semantic_batches,
    semantic_reason_is_too_broad as _semantic_reason_is_too_broad,
    semantic_review_signature as _semantic_review_signature,
    test_ai_connection as _test_ai_connection,
)
from .analyzer import MIN_COMPETITOR_COVERAGE_RATIO, MIN_TARGETING_SEARCH_VOLUME, RULE_ENGINE_VERSION, _clear_generic_decor_mismatch, _has_product_anchor, _is_generic_decor_without_anchor, _negative_phrase_protected_roots, analyze_keyword, derive_negative_phrase_candidates, infer_core_terms, minimum_competitor_coverage, root_candidate_metadata
from .api_support import (
    api_error as _api_error,
    get_product as _get_product,
    keyword_filters as _keyword_filters,
    keyword_response as _keyword_response,
    keyword_update_parts as _keyword_update_parts,
    model_dict as _model_dict,
    normalize_bullets as _normalize_bullets,
    parse_mapping as _parse_mapping,
    product_response as _product_response,
)
from .db import init_db, read_connection, transaction
from .importer import ImportValidationError, _product_dict, _reconcile_broad_root_actions, _update_automatic_analysis, import_parsed_workbook, parse_workbook
from .schemas import AIConfigUpdate, KeywordUpdate, ProductCreate, ProductUpdate, SemanticReviewRequest
from .utils import clean_text, dumps, loads, now_iso, tokens


APP_VERSION = "0.3.2"

# The browser starts reviews as a background job so a refresh does not lose
# the visible progress. Decisions themselves remain durable in SQLite; this
# in-memory registry only tracks the live task and its batch counters.
_REVIEW_JOBS: dict[int, dict[str, Any]] = {}
_REVIEW_JOBS_LOCK = threading.Lock()


app = FastAPI(title="Amazon Keyword Library API", version=APP_VERSION)
cors_origins = [origin.strip() for origin in os.getenv("KEYWORD_CORS_ORIGINS", "http://localhost:3000,http://localhost:5173,http://127.0.0.1:5173").split(",") if origin.strip()]
app.add_middleware(CORSMiddleware, allow_origins=cors_origins, allow_credentials=True, allow_methods=["*"], allow_headers=["*"])


@app.on_event("startup")
def startup() -> None:
    init_db()


@app.get("/health")
@app.get("/api/health")
def health() -> dict[str, Any]:
    try:
        init_db()
        with read_connection() as connection:
            connection.execute("SELECT 1").fetchone()
        db_status = "ok"
    except Exception:
        db_status = "error"
    return {"status": "ok" if db_status == "ok" else "degraded", "service": "amazon-keyword-backend", "database": db_status, "version": APP_VERSION, "timestamp": now_iso()}


@app.get("/api/ai-config")
def get_ai_config() -> dict[str, Any]:
    with read_connection() as connection:
        row = connection.execute("SELECT value_json, updated_at FROM app_settings WHERE key = 'ai_config'").fetchone()
    config = loads(row["value_json"], {}) if row else {}
    return {**_public_ai_config(config), "updated_at": row["updated_at"] if row else None}


@app.put("/api/ai-config")
def update_ai_config(payload: AIConfigUpdate) -> dict[str, Any]:
    values = _model_dict(payload)
    base_url = clean_text(values.get("base_url")).rstrip("/")
    model = clean_text(values.get("model"))
    if not (base_url.startswith("https://") or base_url.startswith("http://")):
        _api_error(422, "invalid_ai_base_url", "AI 接口地址必须以 http:// 或 https:// 开头")
    if not model:
        _api_error(422, "missing_ai_model", "模型名称不能为空")
    with transaction() as connection:
        existing_row = connection.execute("SELECT value_json FROM app_settings WHERE key = 'ai_config'").fetchone()
        existing = loads(existing_row["value_json"], {}) if existing_row else {}
        api_key = clean_text(values.get("api_key")) or clean_text(existing.get("api_key"))
        config = {
            "provider": clean_text(values.get("provider")) or "mimo",
            "base_url": base_url,
            "model": model,
            "api_key": api_key,
            "enabled": bool(values.get("enabled")),
            "timeout_seconds": int(values.get("timeout_seconds") or 60),
        }
        timestamp = now_iso()
        connection.execute(
            """INSERT INTO app_settings(key, value_json, updated_at) VALUES ('ai_config', ?, ?)
               ON CONFLICT(key) DO UPDATE SET value_json = excluded.value_json, updated_at = excluded.updated_at""",
            (dumps(config), timestamp),
        )
    return {**_public_ai_config(config), "updated_at": timestamp}


@app.post("/api/ai-config/test")
def test_ai_config() -> dict[str, Any]:
    return _test_ai_connection(_ai_config_or_error())


def _prepare_full_semantic_review(product_id: int) -> None:
    """Reset unlocked semantic state for an explicit full re-review."""

    with transaction() as connection:
        product = _product_dict(_get_product(connection, product_id))
        _update_automatic_analysis(connection, product_id, product)
        connection.execute(
            """UPDATE keywords
               SET semantic_reviewed = CASE WHEN manual_locked = 1 THEN 1 ELSE 0 END,
                   semantic_reviewed_at = CASE WHEN manual_locked = 1 THEN semantic_reviewed_at ELSE NULL END,
                   semantic_review_signature = CASE WHEN manual_locked = 1 THEN semantic_review_signature ELSE NULL END,
                   updated_at = ?
               WHERE product_id = ? AND deleted_at IS NULL""",
            (now_iso(), product_id),
        )


def _semantic_review_sync(
    product_id: int,
    payload: SemanticReviewRequest,
    progress_callback: Callable[[int, int, bool, str | None], None] | None = None,
    *,
    full_review_prepared: bool = False,
) -> dict[str, Any]:
    """Run local-rule candidates through bounded concurrent semantic batches."""

    config = _ai_config_or_error()
    if payload.review_mode == "full" and not full_review_prepared:
        _prepare_full_semantic_review(product_id)
    with read_connection() as connection:
        product_row = _get_product(connection, product_id)
        competitor_total = int(connection.execute("SELECT COUNT(*) FROM product_asins WHERE product_id = ? AND role = 'competitor'", (product_id,)).fetchone()[0])
        product = {**_product_dict(product_row), "competitor_total": competitor_total}
        source_count_rows = connection.execute("SELECT keyword_id, COUNT(*) AS count FROM keyword_sources WHERE product_id = ? GROUP BY keyword_id", (product_id,)).fetchall()
        keyword_source_counts = {int(item["keyword_id"]): int(item["count"]) for item in source_count_rows}
        if payload.keyword_ids:
            placeholders = ", ".join("?" for _ in payload.keyword_ids)
            locked_clause = " AND manual_locked = 0" if payload.review_mode == "full" else ""
            rows = connection.execute(f"SELECT * FROM keywords WHERE product_id = ? AND deleted_at IS NULL AND id IN ({placeholders}){locked_clause} ORDER BY id ASC LIMIT ?", (product_id, *payload.keyword_ids, payload.limit)).fetchall()
        else:
            rows = connection.execute(
                """SELECT * FROM keywords WHERE product_id = ? AND deleted_at IS NULL AND semantic_reviewed = 0
                   ORDER BY COALESCE(monthly_search_volume, -1) DESC,
                            COALESCE(NULLIF((SELECT COUNT(*) FROM keyword_sources ks WHERE ks.keyword_id = keywords.id), 0), related_product_count, 0) DESC,
                            relevance_score DESC, id ASC LIMIT ?""",
                (product_id, payload.limit),
            ).fetchall()
    if not rows:
        with read_connection() as connection:
            total = connection.execute("SELECT COUNT(*) AS count FROM keywords WHERE product_id = ? AND deleted_at IS NULL", (product_id,)).fetchone()["count"]
        if not total:
            _api_error(422, "no_keywords", "当前产品还没有可供语义审核的关键词")
        promoted = _promote_negative_phrase_recommendations(product_id, product)
        return {"product_id": product_id, "provider": clean_text(config.get("provider")), "model": clean_text(config.get("model")), "review_mode": payload.review_mode, "reviewed": 0, "batches": 0, "already_reviewed": True, "negative_phrase_promoted": promoted, "items": []}

    all_candidates = [
        {"id": int(row["id"]), "keyword": row["keyword_raw"], "translation": row["keyword_translation"], "rule_score": row["relevance_score"], "rule_strength": row["match_strength_auto"], "rule_action": row["suggested_action_auto"], "monthly_search_volume": row["monthly_search_volume"], "source_count": max(0, int(max(int(row["related_product_count"]) if row["related_product_count"] is not None else 0, keyword_source_counts.get(int(row["id"]), 0), len(loads(row["related_asins_json"], []))))), "source_total": competitor_total, "manual_locked": bool(row["manual_locked"]), "manual_action": row["manual_action"], "root_metadata": root_candidate_metadata(row["keyword_raw"], product)}
        for row in rows
    ]
    actions = {"exact", "broad", "negative_exact", "negative_phrase", "observe"}
    applied: list[dict[str, Any]] = []
    core_terms = [clean_text(term).casefold() for term in product.get("core_terms", [])]
    batch_specs = [
        (batch_index, all_candidates[start:start + payload.batch_size])
        for batch_index, start in enumerate(range(0, len(all_candidates), payload.batch_size))
    ]
    batch_results, worker_count = _run_semantic_batches(
        config,
        product,
        batch_specs,
        payload.concurrency,
        _ai_json_response,
    )
    failed_batches: list[dict[str, Any]] = []

    # Network-bound model calls run concurrently above. SQLite writes remain
    # deterministic here, in original batch order, so partial failures never
    # create concurrent database transactions or reorder audit history.
    for batch_index, candidates in batch_specs:
        reviews_by_id, batch_error = batch_results.get(batch_index, ({}, "批次没有返回结果"))
        if batch_error:
            failed_batches.append({"batch": batch_index + 1, "count": len(candidates), "error": batch_error})
            if progress_callback:
                progress_callback(batch_index + 1, len(candidates), False, batch_error)
            continue

        timestamp = now_iso()
        with transaction() as connection:
            for candidate in candidates:
                review = reviews_by_id[candidate["id"]]
                keyword_id = candidate["id"]
                decision = clean_text(review.get("decision")).casefold()
                if decision not in actions:
                    decision = "observe"
                keyword_text = clean_text(candidate["keyword"]).casefold()
                score = max(0, min(100, int(review.get("relevance_score", candidate["rule_score"]) or 0)))
                confidence = max(0.0, min(1.0, float(review.get("confidence", 0.5) or 0.5)))
                reason = clean_text(review.get("reason_zh"))[:600] or "AI 语义审核未给出详细理由"
                rule_adjusted = False
                if decision in {"exact", "broad"} and candidate["monthly_search_volume"] is not None and int(candidate["monthly_search_volume"]) < MIN_TARGETING_SEARCH_VOLUME:
                    decision = "observe"
                    reason += f"；月搜索量仅 {int(candidate['monthly_search_volume'])}，低于 {MIN_TARGETING_SEARCH_VOLUME} 投放门槛，降为观察"
                    rule_adjusted = True
                if decision in {"exact", "broad"} and not candidate["source_total"]:
                    decision = "observe"
                    reason += "；竞品 ASIN 覆盖数据尚未导入，无法计算相关性占比，降为观察"
                    rule_adjusted = True
                else:
                    minimum_coverage = minimum_competitor_coverage(candidate["source_total"])
                    if decision in {"exact", "broad"} and candidate["source_count"] < minimum_coverage:
                        decision = "observe"
                        reason += f"；竞品覆盖仅 {candidate['source_count']}/{candidate['source_total']}，低于 {minimum_coverage}/{candidate['source_total']}（{MIN_COMPETITOR_COVERAGE_RATIO:.0%}相关性门槛），降为观察"
                        rule_adjusted = True
                # Keep a clear, low-scoring generic-decor mismatch as a
                # negative-exact draft even if the model answers observe or
                # exact. This is the double-audit guard for cases such as
                # `room decor`: the complete query is excluded, while the
                # single root `room` remains usable in a wreath combination.
                if candidate["rule_action"] == "negative_exact" and decision in {"observe", "exact", "broad"} and score < 50:
                    if _clear_generic_decor_mismatch(keyword_text, analyze_keyword(keyword_text, product), product):
                        decision = "negative_exact"
                        reason += "；内置规则与 AI 均判定为泛房间装饰查询，保留否定精准"
                        rule_adjusted = True
                if decision in {"exact", "broad"} and keyword_text not in core_terms and _is_generic_decor_without_anchor(keyword_text, product):
                    decision = "observe"
                    reason += "；装饰类意图过宽且缺少花环/产品类型锚点，降为观察/暂不投放"
                    rule_adjusted = True
                # The semantic model can describe a query as too broad while accidentally
                # returning `exact`. For a non-core short generic query,
                # honour that semantic warning and keep the term out of the
                # exact budget.
                if decision in {"exact", "broad"} and keyword_text not in core_terms:
                    if _is_short_generic_query(keyword_text, product) and (_semantic_reason_is_too_broad(reason) or score < 80):
                        decision = "observe"
                        reason += "；语义审核提示词义过宽或具体度不足，降为观察/暂不投放"
                        rule_adjusted = True
                # A complete, sufficiently searched product root is the one
                # intentional broad seed. Long-tail terms are never promoted
                # by this guard.
                if candidate["rule_action"] == "broad" and decision == "exact" and score >= 60:
                    decision = "broad"
                    reason += "；完整核心词根且语义相关度达标，保留广泛抓词根建议"
                    rule_adjusted = True
                if decision == "broad" and keyword_text not in core_terms:
                    decision = "exact"
                    reason += "；广泛仅允许使用完整核心词根，已降为精准草稿"
                    rule_adjusted = True
                if decision == "negative_phrase":
                    root = clean_text(review.get("negative_phrase_root") or keyword_text).casefold()
                    related = connection.execute("SELECT relevance_score FROM keywords WHERE product_id = ? AND deleted_at IS NULL AND keyword_normalized LIKE ?", (product_id, f"%{root}%")).fetchall()
                    # Keep the complete generic-decor example at negative
                    # exact. `room decor` itself can be excluded, but a
                    # phrase-level action would be too aggressive because
                    # valid combinations such as `room ... wreath` may still
                    # exist in the same corpus.
                    if _clear_generic_decor_mismatch(keyword_text, analyze_keyword(keyword_text, product), product):
                        decision = "negative_exact"
                        reason += "；泛房间装饰只否定完整搜索词，避免把 room 相关有效组合一并拦截"
                    # A one-word root such as `room` is too broad to negate as
                    # a phrase. Require a multi-word root plus repeated-root /
                    # relevance conflict checks.
                    if decision == "negative_phrase" and (len(tokens(root)) < 2 or len(related) < 2 or any(int(item["relevance_score"] or 0) >= 50 for item in related)):
                        decision = "negative_exact"
                        reason += "；否定词组可能误伤相关词或词根过短，已降为否定精准草稿"
                        rule_adjusted = True
                strength = "strong" if score >= 80 else "medium" if score >= 50 else "weak" if score >= 20 else "irrelevant"
                signature = _semantic_review_signature(product, candidate)
                final_source = "manual" if candidate.get("manual_action") else "rule_gate" if rule_adjusted else "semantic"
                conflict_actions = []
                if candidate.get("manual_action") and candidate["manual_action"] != decision:
                    conflict_actions = sorted({clean_text(candidate["manual_action"]).casefold(), decision})
                if not candidate["manual_locked"]:
                    connection.execute("""UPDATE keywords SET relevance_score = ?, match_strength_auto = ?, suggested_action_auto = ?, suggested_match_type = ?,
                        advice_reason = ?, advice_confidence = ?, advice_risk_level = ?, semantic_reviewed = 1, semantic_reviewed_at = ?, semantic_review_signature = ?,
                        rule_engine_version = ?, rule_action_before_semantic = ?, final_action_source = ?, conflict_actions_json = ?,
                        protected_terms_json = ?, updated_at = ? WHERE id = ?""",
                        (score, strength, decision, decision, "AI 语义审核：" + reason, confidence, "high" if decision == "negative_phrase" else "medium" if decision in {"broad", "negative_exact", "observe"} else "low", timestamp, signature, RULE_ENGINE_VERSION, candidate["rule_action"], final_source, dumps(conflict_actions), dumps(sorted(_negative_phrase_protected_roots(product))), timestamp, keyword_id),
                    )
                else:
                    connection.execute("UPDATE keywords SET semantic_reviewed = 1, semantic_reviewed_at = ?, semantic_review_signature = ?, rule_engine_version = ?, rule_action_before_semantic = ?, final_action_source = ?, conflict_actions_json = ?, updated_at = ? WHERE id = ?", (timestamp, signature, RULE_ENGINE_VERSION, candidate["rule_action"], final_source, dumps(conflict_actions), timestamp, keyword_id))
                connection.execute("INSERT INTO audit_logs(product_id, keyword_id, action, details_json, created_at) VALUES (?, ?, 'mimo_semantic_review', ?, ?)", (product_id, keyword_id, dumps({"decision": decision, "score": score, "confidence": confidence, "manual_locked": candidate["manual_locked"], "rule_action_before_semantic": candidate["rule_action"], "final_action_source": final_source, "conflict_actions": conflict_actions}), timestamp))
                applied.append({"id": keyword_id, "keyword": candidate["keyword"], "decision": decision, "relevance_score": score, "reason": reason, "manual_locked": candidate["manual_locked"], "final_action_source": final_source, "conflict_actions": conflict_actions})
        if progress_callback:
            progress_callback(batch_index + 1, len(candidates), True, None)

    if failed_batches and not applied:
        first_error = failed_batches[0]["error"]
        _api_error(502, "ai_review_failed", f"AI 语义审核失败，{len(failed_batches)} 批均未完成：{first_error}")

    # Phrase negatives are derived only after the entire product-level set of
    # exact decisions has been persisted.  Broad roots are reconciled at the
    # same boundary so the product-level ten-term cap is deterministic.
    promoted = _promote_negative_phrase_recommendations(product_id, product)
    with transaction() as connection:
        ranked_broad = _reconcile_broad_root_actions(connection, product_id, product)
    promoted_by_id = {item["id"]: item for item in promoted}
    for item in applied:
        if item["id"] in promoted_by_id:
            item.update({"decision": "negative_phrase", "negative_phrase_root": promoted_by_id[item["id"]]["root"], "reason": promoted_by_id[item["id"]]["reason"]})

    return {
        "product_id": product_id,
        "provider": clean_text(config.get("provider")),
        "model": clean_text(config.get("model")),
        "review_mode": payload.review_mode,
        "reviewed": len(applied),
        "batches": len(batch_specs),
        "successful_batches": len(batch_specs) - len(failed_batches),
        "failed_batches": failed_batches,
        "partial": bool(failed_batches),
        "concurrency": worker_count,
        "negative_phrase_promoted": promoted,
        "broad_root_candidates": ranked_broad,
        "items": applied,
    }


def _promote_negative_phrase_recommendations(product_id: int, product: dict[str, Any]) -> list[dict[str, Any]]:
    """Persist safe product-level phrase-negative drafts after semantic review."""

    with transaction() as connection:
        rows = [dict(row) for row in connection.execute("SELECT * FROM keywords WHERE product_id = ? AND deleted_at IS NULL", (product_id,)).fetchall()]
        candidates = derive_negative_phrase_candidates(rows, product)
        promoted: list[dict[str, Any]] = []
        timestamp = now_iso()
        promoted_ids: set[int] = set()
        for root, evidence in sorted(candidates.items(), key=lambda item: (-int(item[1].get("root_length", len(item[0].split()))), item[0])):
            representative = connection.execute(
                "SELECT id, keyword_raw, manual_locked, manual_action, suggested_action_auto FROM keywords WHERE id = ? AND product_id = ? AND deleted_at IS NULL",
                (evidence["representative_id"], product_id),
            ).fetchone()
            if representative is None or int(representative["id"]) in promoted_ids or representative["manual_locked"] or representative["manual_action"]:
                continue
            if representative["suggested_action_auto"] != "negative_exact":
                continue
            reason = f"基于否定精准词根二次审核：建议否定词组“{root}”；已审核 {evidence['affected_count']} 个受影响词，未发现相关产品组合"
            impact = {
                "root": root,
                "seed_count": evidence["seed_count"],
                "seed_keywords": evidence["seed_keywords"],
                "affected_count": evidence["affected_count"],
                "affected_keywords": evidence["affected_keywords"][:50],
                "protected_keywords": evidence["protected_keywords"],
                "blocked": False,
                "basis": "已审核否定精准词的词根重复度 + 受影响词语义结果 + 产品锚点误伤检查",
            }
            connection.execute(
                """UPDATE keywords SET suggested_action_auto = 'negative_phrase', suggested_match_type = 'negative_phrase',
                   advice_reason = ?, advice_risk_level = 'high', negative_impact_json = ?, negative_phrase_eligible = 1,
                   negative_phrase_evidence_json = ?, final_action_source = 'rule_gate', updated_at = ? WHERE id = ?""",
                (reason, dumps(impact), dumps(evidence), timestamp, representative["id"]),
            )
            connection.execute(
                "INSERT INTO audit_logs(product_id, keyword_id, action, details_json, created_at) VALUES (?, ?, 'negative_phrase_derived', ?, ?)",
                (product_id, representative["id"], dumps(impact), timestamp),
            )
            promoted_ids.add(int(representative["id"]))
            promoted.append({"id": int(representative["id"]), "keyword": representative["keyword_raw"], "root": root, "affected_count": evidence["affected_count"], "reason": reason})
        return promoted


def _semantic_review_counts(product_id: int) -> tuple[int, int]:
    with read_connection() as connection:
        row = connection.execute(
            """SELECT COUNT(*) AS total,
                      SUM(CASE WHEN semantic_reviewed = 1 THEN 1 ELSE 0 END) AS reviewed
               FROM keywords WHERE product_id = ? AND deleted_at IS NULL""",
            (product_id,),
        ).fetchone()
    return int(row["total"] or 0), int(row["reviewed"] or 0)


def _semantic_review_status_snapshot(product_id: int) -> dict[str, Any]:
    total, reviewed = _semantic_review_counts(product_id)
    with _REVIEW_JOBS_LOCK:
        state = dict(_REVIEW_JOBS.get(product_id) or {})
    status = state.get("status", "idle")
    if status == "completed" and total > reviewed:
        status = "idle"
    return {
        "product_id": product_id,
        "status": status,
        "reviewed": reviewed,
        "total": total,
        "pending": max(0, total - reviewed),
        "batches": int(state.get("batches_total") or 0),
        "batches_total": int(state.get("batches_total") or 0),
        "batches_completed": int(state.get("batches_completed") or 0),
        "successful_batches": int(state.get("successful_batches") or 0),
        "failed_batches": state.get("failed_batches") or [],
        "started_at": state.get("started_at"),
        "updated_at": state.get("updated_at"),
        "completed_at": state.get("completed_at"),
        "error": state.get("error"),
        "review_mode": state.get("review_mode", "incremental"),
        "items": [],
    }


def _run_background_semantic_review(product_id: int, payload: SemanticReviewRequest) -> None:
    def on_batch_complete(batch_number: int, count: int, success: bool, error: str | None) -> None:
        with _REVIEW_JOBS_LOCK:
            state = _REVIEW_JOBS.get(product_id)
            if state is None:
                return
            state["batches_completed"] = max(int(state.get("batches_completed") or 0), batch_number)
            if success:
                state["successful_batches"] = int(state.get("successful_batches") or 0) + 1
            elif error:
                state.setdefault("failed_batches", []).append({"batch": batch_number, "count": count, "error": error})
            state["updated_at"] = now_iso()

    try:
        result = _semantic_review_sync(product_id, payload, progress_callback=on_batch_complete, full_review_prepared=payload.review_mode == "full")
        with _REVIEW_JOBS_LOCK:
            state = _REVIEW_JOBS.get(product_id)
            if state is not None:
                state["status"] = "partial" if result.get("partial") else "completed"
                state["failed_batches"] = result.get("failed_batches") or []
                state["successful_batches"] = int(result.get("successful_batches") or 0)
                state["error"] = (result.get("failed_batches") or [{}])[0].get("error") if result.get("partial") else None
                state["completed_at"] = now_iso()
                state["updated_at"] = state["completed_at"]
    except Exception as error:
        with _REVIEW_JOBS_LOCK:
            state = _REVIEW_JOBS.get(product_id)
            if state is not None:
                state["status"] = "failed"
                state["error"] = str(error)
                state["completed_at"] = now_iso()
                state["updated_at"] = state["completed_at"]


def _start_background_semantic_review(product_id: int, payload: SemanticReviewRequest) -> dict[str, Any]:
    with read_connection() as connection:
        _get_product(connection, product_id)
    with _REVIEW_JOBS_LOCK:
        current = _REVIEW_JOBS.get(product_id)
        already_running = bool(current and current.get("status") == "running")
    if already_running:
        return _semantic_review_status_snapshot(product_id)
    if payload.review_mode == "full":
        _prepare_full_semantic_review(product_id)
    total, reviewed = _semantic_review_counts(product_id)
    if not total:
        _api_error(422, "no_keywords", "当前产品还没有可供语义审核的关键词")
    pending = max(0, total - reviewed)
    if not pending:
        return _semantic_review_status_snapshot(product_id)
    candidate_count = min(pending, payload.limit)
    batch_total = math.ceil(candidate_count / payload.batch_size)
    already_running = False
    with _REVIEW_JOBS_LOCK:
        current = _REVIEW_JOBS.get(product_id)
        if current and current.get("status") == "running":
            already_running = True
        else:
            timestamp = now_iso()
            _REVIEW_JOBS[product_id] = {
                "status": "running", "batches_total": batch_total, "batches_completed": 0,
                "successful_batches": 0, "failed_batches": [], "started_at": timestamp,
                "updated_at": timestamp, "completed_at": None, "error": None, "review_mode": payload.review_mode,
            }
    if already_running:
        return _semantic_review_status_snapshot(product_id)
    thread = threading.Thread(target=_run_background_semantic_review, args=(product_id, payload), name=f"ai-review-{product_id}", daemon=True)
    thread.start()
    return _semantic_review_status_snapshot(product_id)


@app.post("/api/products/{product_id}/semantic-review")
def semantic_review(product_id: int, payload: SemanticReviewRequest) -> dict[str, Any]:
    """Start refresh-safe background review or run the synchronous API mode."""

    if payload.background:
        _ai_config_or_error()
        return _start_background_semantic_review(product_id, payload)
    return _semantic_review_sync(product_id, payload)


@app.get("/api/products/{product_id}/semantic-review/status")
def semantic_review_status(product_id: int) -> dict[str, Any]:
    with read_connection() as connection:
        _get_product(connection, product_id)
    return _semantic_review_status_snapshot(product_id)


@app.get("/api/products")
def list_products(
    search: str | None = None,
    site: str | None = None,
    status: str | None = None,
    include_deleted: bool = False,
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=200),
) -> dict[str, Any]:
    conditions = ["1=1"]
    params: list[Any] = []
    if not include_deleted:
        conditions.append("deleted_at IS NULL")
    if search:
        conditions.append("(LOWER(name) LIKE ? OR LOWER(COALESCE(asin, '')) LIKE ?)")
        term = f"%{clean_text(search).casefold()}%"
        params.extend([term, term])
    if site:
        conditions.append("site = ?")
        params.append(site)
    if status:
        conditions.append("status = ?")
        params.append(status)
    where = " AND ".join(conditions)
    with read_connection() as connection:
        total = int(connection.execute(f"SELECT COUNT(*) FROM products WHERE {where}", params).fetchone()[0])
        rows = connection.execute(f"SELECT * FROM products WHERE {where} ORDER BY updated_at DESC, id DESC LIMIT ? OFFSET ?", [*params, page_size, (page - 1) * page_size]).fetchall()
        return {"items": [_product_response(row, connection) for row in rows], "page": page, "page_size": page_size, "total": total, "pages": math.ceil(total / page_size) if total else 0}


@app.post("/api/products", status_code=201)
def create_product(payload: ProductCreate) -> dict[str, Any]:
    values = _model_dict(payload)
    name = clean_text(values.get("name"))
    site = clean_text(values.get("site"))
    if not name or not site:
        _api_error(422, "missing_product_fields", "产品名称和站点不能为空")
    bullets = _normalize_bullets(values.get("bullet_points") if values.get("bullet_points") is not None else values.get("five_points"))
    title = clean_text(values.get("product_title") or values.get("title")) or None
    description = clean_text(values.get("product_description") or values.get("description")) or None
    core_terms = values.get("core_terms") or infer_core_terms({"product_title": title, "bullet_points": bullets})
    timestamp = now_iso()
    with transaction() as connection:
        cursor = connection.execute(
            """INSERT INTO products(name, asin, site, language, brand, category, product_title, bullet_points_json,
               product_description, search_terms, core_terms_json, excluded_terms_json, settings_json, status,
               created_at, updated_at) VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)""",
            (name, clean_text(values.get("asin")) or None, site, clean_text(values.get("language")) or "en_US", clean_text(values.get("brand")) or None,
             clean_text(values.get("category")) or None, title, dumps(bullets), description, clean_text(values.get("search_terms")) or None,
             dumps(core_terms), dumps(values.get("excluded_terms") or []), dumps(values.get("settings") or {}), values.get("status") or "preparing", timestamp, timestamp),
        )
        product_id = int(cursor.lastrowid)
        if clean_text(values.get("asin")):
            asin = clean_text(values["asin"]).upper()
            connection.execute("INSERT INTO product_asins(product_id, asin, role, first_seen_at, last_seen_at) VALUES (?, ?, 'own', ?, ?)", (product_id, asin, timestamp, timestamp))
        row = connection.execute("SELECT * FROM products WHERE id = ?", (product_id,)).fetchone()
        return _product_response(row, connection)


@app.get("/api/products/{product_id}")
def get_product(product_id: int) -> dict[str, Any]:
    with read_connection() as connection:
        row = _get_product(connection, product_id)
        return _product_response(row, connection)


@app.patch("/api/products/{product_id}")
def update_product(product_id: int, payload: ProductUpdate) -> dict[str, Any]:
    values = _model_dict(payload, exclude_unset=True)
    timestamp = now_iso()
    with transaction() as connection:
        row = _get_product(connection, product_id)
        if "site" in values and clean_text(values["site"]) != row["site"]:
            count = connection.execute("SELECT COUNT(*) FROM keywords WHERE product_id = ? AND deleted_at IS NULL", (product_id,)).fetchone()[0]
            if count:
                _api_error(409, "site_change_blocked", "已有关键词的产品不能直接更换站点")
        assignments: list[str] = []
        params: list[Any] = []
        aliases = {
            "product_title": "product_title", "title": "product_title", "product_description": "product_description", "description": "product_description",
            "bullet_points": "bullet_points_json", "five_points": "bullet_points_json", "core_terms": "core_terms_json", "excluded_terms": "excluded_terms_json", "settings": "settings_json",
        }
        for key, value in values.items():
            target = aliases.get(key, key)
            if target == "bullet_points_json":
                if "bullet_points" in values:
                    if key != "bullet_points":
                        continue
                    value = dumps(_normalize_bullets(value))
                else:
                    value = dumps(_normalize_bullets(value))
            elif target in {"core_terms_json", "excluded_terms_json", "settings_json"}:
                value = dumps(value or ([] if target != "settings_json" else {}))
            elif target in {"product_title", "product_description", "name", "site", "language", "asin", "brand", "category", "search_terms"}:
                value = clean_text(value) or None
            if target not in {"title", "description", "five_points"}:
                assignments.append(f"{target} = ?")
                params.append(value)
        if assignments:
            assignments.extend(["updated_at = ?"])
            params.extend([timestamp, product_id])
            connection.execute(f"UPDATE products SET {', '.join(assignments)} WHERE id = ?", params)
        updated = connection.execute("SELECT * FROM products WHERE id = ?", (product_id,)).fetchone()
        if any(key in values for key in ("product_title", "title", "bullet_points", "five_points", "product_description", "description", "search_terms", "core_terms", "excluded_terms", "brand")):
            _update_automatic_analysis(connection, product_id, _product_dict(updated))
            # Product-copy or root changes invalidate prior semantic decisions;
            # the next audit must run against the new evidence.
            connection.execute("UPDATE keywords SET semantic_reviewed = 0, semantic_reviewed_at = NULL, semantic_review_signature = NULL WHERE product_id = ? AND deleted_at IS NULL", (product_id,))
        connection.execute("INSERT INTO audit_logs(product_id, action, details_json, created_at) VALUES (?, 'product_update', ?, ?)", (product_id, dumps(values), timestamp))
        return _product_response(updated, connection)


@app.delete("/api/products/{product_id}")
def archive_product(product_id: int) -> dict[str, Any]:
    timestamp = now_iso()
    with transaction() as connection:
        _get_product(connection, product_id)
        connection.execute("UPDATE products SET status = 'archived', deleted_at = ?, updated_at = ? WHERE id = ?", (timestamp, timestamp, product_id))
        connection.execute("INSERT INTO audit_logs(product_id, action, details_json, created_at) VALUES (?, 'product_archive', '{}', ?)", (product_id, timestamp))
        return {"id": product_id, "status": "archived", "deleted_at": timestamp}


@app.post("/api/products/{product_id}/restore")
def restore_product(product_id: int) -> dict[str, Any]:
    timestamp = now_iso()
    with transaction() as connection:
        row = _get_product(connection, product_id, include_deleted=True)
        connection.execute("UPDATE products SET status = 'preparing', deleted_at = NULL, updated_at = ? WHERE id = ?", (timestamp, product_id))
        return _product_response(connection.execute("SELECT * FROM products WHERE id = ?", (row["id"],)).fetchone(), connection)


@app.post("/api/products/{product_id}/imports")
@app.post("/api/products/{product_id}/import")
async def import_keywords(
    product_id: int,
    file: UploadFile = File(...),
    sheet_name: str | None = Form(None),
    mapping_json: str | None = Form(None),
) -> dict[str, Any]:
    data = await file.read()
    filename = file.filename or "uploaded-file"
    mapping = _parse_mapping(mapping_json)
    try:
        with read_connection() as connection:
            product = _get_product(connection, product_id)
        parsed = parse_workbook(data, filename, sheet_name=sheet_name, supplied_mapping=mapping)
        result = import_parsed_workbook(product_id, filename, parsed, product)
        return result
    except ImportValidationError as exc:
        _api_error(400, exc.code, str(exc))
    except HTTPException:
        raise
    except Exception as exc:
        _api_error(500, "import_failed", f"导入失败: {exc}")


@app.get("/api/products/{product_id}/imports")
def list_imports(product_id: int, page: int = Query(1, ge=1), page_size: int = Query(50, ge=1, le=200)) -> dict[str, Any]:
    with read_connection() as connection:
        _get_product(connection, product_id)
        total = connection.execute("SELECT COUNT(*) FROM imports WHERE product_id = ?", (product_id,)).fetchone()[0]
        rows = connection.execute("SELECT * FROM imports WHERE product_id = ? ORDER BY created_at DESC, id DESC LIMIT ? OFFSET ?", (product_id, page_size, (page - 1) * page_size)).fetchall()
        items = []
        for row in rows:
            item = dict(row)
            for key in ("error_details_json", "unmapped_headers_json", "source_asins_json", "mapping_json"):
                item[key.removesuffix("_json")] = loads(item.pop(key, "[]" if key != "mapping_json" else "{}"), [] if key != "mapping_json" else {})
            items.append(item)
        return {"items": items, "page": page, "page_size": page_size, "total": int(total), "pages": math.ceil(total / page_size) if total else 0}


@app.get("/api/imports/{import_id}")
def get_import(import_id: int) -> dict[str, Any]:
    with read_connection() as connection:
        row = connection.execute("SELECT * FROM imports WHERE id = ?", (import_id,)).fetchone()
        if row is None:
            _api_error(404, "import_not_found", "导入记录不存在")
        item = dict(row)
        for key in ("error_details_json", "unmapped_headers_json", "source_asins_json", "mapping_json"):
            item[key.removesuffix("_json")] = loads(item.pop(key, "[]" if key != "mapping_json" else "{}"), [] if key != "mapping_json" else {})
        return item


@app.get("/api/products/{product_id}/keywords")
def list_keywords(
    product_id: int,
    search: str | None = None,
    q: str | None = None,
    include_root: str | None = None,
    exclude_root: str | None = None,
    match_strength: str | None = None,
    category: str | None = None,
    traffic_type: str | None = None,
    source_asin: str | None = None,
    action: str | None = None,
    suggested_action: str | None = None,
    status: str | None = None,
    min_search_volume: int | None = Query(None, ge=0),
    max_search_volume: int | None = Query(None, ge=0),
    sort_by: str = "monthly_search_volume",
    sort_order: str = Query("desc", pattern="^(asc|desc)$"),
    include_deleted: bool = False,
    page: int = Query(1, ge=1),
    page_size: int = Query(50, ge=1, le=500),
) -> dict[str, Any]:
    if max_search_volume is not None and min_search_volume is not None and max_search_volume < min_search_volume:
        _api_error(422, "invalid_range", "搜索量范围无效")
    selected_action = action or suggested_action
    sort_map = {
        "relevance": "k.relevance_score", "relevance_score": "k.relevance_score", "monthly_search_volume": "k.monthly_search_volume",
        "aba_weekly_rank": "k.aba_weekly_rank", "traffic_share": "k.traffic_share", "ppc_bid": "k.ppc_bid", "updated_at": "k.updated_at", "keyword": "k.keyword_normalized",
    }
    sort_expression = sort_map.get(sort_by, sort_map["monthly_search_volume"])
    where, params = _keyword_filters(product_id, search=search or q, include_root=include_root, exclude_root=exclude_root, match_strength=match_strength, category=category, traffic_type=traffic_type, source_asin=source_asin, action=selected_action, status=status, min_search_volume=min_search_volume, max_search_volume=max_search_volume, include_deleted=include_deleted)
    with read_connection() as connection:
        _get_product(connection, product_id)
        total = int(connection.execute(f"SELECT COUNT(*) FROM keywords k WHERE {where}", params).fetchone()[0])
        rows = connection.execute(f"SELECT k.* FROM keywords k WHERE {where} ORDER BY {sort_expression} {sort_order.upper()}, k.id ASC LIMIT ? OFFSET ?", [*params, page_size, (page - 1) * page_size]).fetchall()
        return {"items": [_keyword_response(row, connection) for row in rows], "page": page, "page_size": page_size, "total": total, "pages": math.ceil(total / page_size) if total else 0}


# Register this static path before the dynamic ``{keyword_id}`` route below;
# otherwise a request for ``/keywords/export`` would be interpreted as a
# keyword id and return a 422 validation error.
@app.get("/api/products/{product_id}/keywords/export")
def export_keywords_static(product_id: int, format: str = Query("csv", pattern="^(csv|xlsx)$"), search: str | None = None, match_strength: str | None = None, action: str | None = None) -> StreamingResponse:
    return export_keywords(product_id, format=format, search=search, match_strength=match_strength, action=action)


@app.get("/api/products/{product_id}/keywords/{keyword_id}")
def get_keyword(product_id: int, keyword_id: int) -> dict[str, Any]:
    with read_connection() as connection:
        _get_product(connection, product_id)
        row = connection.execute("SELECT * FROM keywords WHERE id = ? AND product_id = ?", (keyword_id, product_id)).fetchone()
        if row is None:
            _api_error(404, "keyword_not_found", "关键词不存在")
        data = _keyword_response(row, connection)
        data["history"] = [dict(item) | {"snapshot": loads(item["snapshot_json"], {})} for item in connection.execute("SELECT id, import_id, snapshot_json, captured_at FROM keyword_metric_history WHERE keyword_id = ? ORDER BY captured_at DESC, id DESC LIMIT 100", (keyword_id,)).fetchall()]
        for item in data["history"]:
            item.pop("snapshot_json", None)
        return data


@app.patch("/api/products/{product_id}/keywords/{keyword_id}")
def update_keyword(product_id: int, keyword_id: int, payload: KeywordUpdate) -> dict[str, Any]:
    values = _model_dict(payload, exclude_unset=True)
    assignments, params = _keyword_update_parts(values)
    timestamp = now_iso()
    with transaction() as connection:
        _get_product(connection, product_id)
        row = connection.execute("SELECT * FROM keywords WHERE id = ? AND product_id = ?", (keyword_id, product_id)).fetchone()
        if row is None:
            _api_error(404, "keyword_not_found", "关键词不存在")
        action_supplied = "manual_action" in values or "action" in values
        manual_action = values.get("manual_action", values.get("action"))
        if action_supplied:
            normalized_manual = clean_text(manual_action).casefold()
            normalized_auto = clean_text(row["suggested_action_auto"]).casefold()
            conflict = sorted({normalized_manual, normalized_auto}) if normalized_manual and normalized_manual != normalized_auto else []
            assignments.extend(["conflict_actions_json = ?", "final_action_source = ?"])
            params.extend([dumps(conflict), "manual" if normalized_manual else "rule"])
        if assignments:
            assignments.append("updated_at = ?")
            params.extend([timestamp, keyword_id])
            connection.execute(f"UPDATE keywords SET {', '.join(assignments)} WHERE id = ?", params)
        connection.execute("INSERT INTO audit_logs(product_id, keyword_id, action, details_json, created_at) VALUES (?, ?, 'keyword_manual_update', ?, ?)", (product_id, keyword_id, dumps(values), timestamp))
        updated = connection.execute("SELECT * FROM keywords WHERE id = ?", (keyword_id,)).fetchone()
        return _keyword_response(updated, connection)


@app.post("/api/products/{product_id}/keywords/bulk-update")
def bulk_update_keywords(product_id: int, payload: dict[str, Any] = Body(...)) -> dict[str, Any]:
    keyword_ids = payload.get("keyword_ids") or payload.get("ids") or []
    if not isinstance(keyword_ids, list) or not keyword_ids:
        _api_error(422, "missing_keyword_ids", "keyword_ids 不能为空")
    values = {key: value for key, value in payload.items() if key not in {"keyword_ids", "ids"}}
    update_values = _model_dict(KeywordUpdate(**values), exclude_unset=True)
    assignments, update_params = _keyword_update_parts(update_values)
    updated = 0
    timestamp = now_iso()
    with transaction() as connection:
        _get_product(connection, product_id)
        for keyword_id in keyword_ids:
            current = connection.execute("SELECT id FROM keywords WHERE id = ? AND product_id = ?", (keyword_id, product_id)).fetchone()
            if current is None:
                continue
            if assignments:
                connection.execute(f"UPDATE keywords SET {', '.join(assignments)}, updated_at = ? WHERE id = ? AND product_id = ?", (*update_params, timestamp, int(keyword_id), product_id))
            connection.execute("INSERT INTO audit_logs(product_id, keyword_id, action, details_json, created_at) VALUES (?, ?, 'keyword_manual_update', ?, ?)", (product_id, int(keyword_id), dumps(update_values), timestamp))
            updated += 1
    return {"updated": updated, "keyword_ids": keyword_ids}


@app.delete("/api/products/{product_id}/keywords/{keyword_id}")
def archive_keyword(product_id: int, keyword_id: int) -> dict[str, Any]:
    timestamp = now_iso()
    with transaction() as connection:
        _get_product(connection, product_id)
        row = connection.execute("SELECT id FROM keywords WHERE id = ? AND product_id = ?", (keyword_id, product_id)).fetchone()
        if row is None:
            _api_error(404, "keyword_not_found", "关键词不存在")
        connection.execute("UPDATE keywords SET status = 'deleted', deleted_at = ?, updated_at = ? WHERE id = ?", (timestamp, timestamp, keyword_id))
        return {"id": keyword_id, "status": "deleted", "deleted_at": timestamp}


@app.post("/api/products/{product_id}/keywords/{keyword_id}/restore")
def restore_keyword(product_id: int, keyword_id: int) -> dict[str, Any]:
    timestamp = now_iso()
    with transaction() as connection:
        _get_product(connection, product_id, include_deleted=True)
        row = connection.execute("SELECT * FROM keywords WHERE id = ? AND product_id = ?", (keyword_id, product_id)).fetchone()
        if row is None:
            _api_error(404, "keyword_not_found", "关键词不存在")
        connection.execute("UPDATE keywords SET status = 'active', deleted_at = NULL, updated_at = ? WHERE id = ?", (timestamp, keyword_id))
        return _keyword_response(connection.execute("SELECT * FROM keywords WHERE id = ?", (keyword_id,)).fetchone(), connection)


@app.get("/api/products/{product_id}/stats")
@app.get("/api/products/{product_id}/summary")
def product_stats(product_id: int) -> dict[str, Any]:
    with read_connection() as connection:
        _get_product(connection, product_id)
        total = int(connection.execute("SELECT COUNT(*) FROM keywords WHERE product_id = ? AND deleted_at IS NULL", (product_id,)).fetchone()[0])
        strength_rows = connection.execute("SELECT COALESCE(manual_match_strength, match_strength_auto) AS value, COUNT(*) AS count FROM keywords WHERE product_id = ? AND deleted_at IS NULL GROUP BY value", (product_id,)).fetchall()
        category_rows = connection.execute("SELECT COALESCE(manual_category, category_auto) AS value, COUNT(*) AS count FROM keywords WHERE product_id = ? AND deleted_at IS NULL GROUP BY value ORDER BY count DESC", (product_id,)).fetchall()
        action_rows = connection.execute("SELECT COALESCE(manual_action, suggested_action_auto) AS value, COUNT(*) AS count FROM keywords WHERE product_id = ? AND deleted_at IS NULL GROUP BY value ORDER BY count DESC", (product_id,)).fetchall()
        conflict_count = int(connection.execute("SELECT COUNT(*) FROM keywords WHERE product_id = ? AND deleted_at IS NULL AND conflict_actions_json NOT IN ('[]', '')", (product_id,)).fetchone()[0])
        source_count = int(connection.execute("SELECT COUNT(*) FROM product_asins WHERE product_id = ?", (product_id,)).fetchone()[0])
        return {"product_id": product_id, "total_keywords": total, "source_asins": source_count, "conflict_count": conflict_count, "by_match_strength": {row["value"]: int(row["count"]) for row in strength_rows if row["value"]}, "by_category": {row["value"]: int(row["count"]) for row in category_rows if row["value"]}, "by_suggested_action": {row["value"]: int(row["count"]) for row in action_rows if row["value"]}, "updated_at": now_iso()}


@app.get("/api/products/{product_id}/asins")
def list_product_asins(product_id: int) -> dict[str, Any]:
    with read_connection() as connection:
        _get_product(connection, product_id)
        rows = connection.execute("SELECT * FROM product_asins WHERE product_id = ? ORDER BY role, asin", (product_id,)).fetchall()
        return {"items": [dict(row) for row in rows], "total": len(rows)}


@app.get("/api/products/{product_id}/keywords/export")
def export_keywords(product_id: int, format: str = Query("csv", pattern="^(csv|xlsx)$"), search: str | None = None, match_strength: str | None = None, action: str | None = None) -> StreamingResponse:
    where, params = _keyword_filters(product_id, search=search, include_root=None, exclude_root=None, match_strength=match_strength, category=None, traffic_type=None, source_asin=None, action=action, status=None, min_search_volume=None, max_search_volume=None, include_deleted=False)
    with read_connection() as connection:
        _get_product(connection, product_id)
        rows = connection.execute(f"SELECT k.* FROM keywords k WHERE {where} ORDER BY COALESCE(k.monthly_search_volume, -1) DESC, k.id ASC", params).fetchall()
        output = io.BytesIO()
        headers = ["关键词", "标准化关键词", "相关性分", "竞品覆盖", "匹配强度", "建议动作", "最终动作来源", "规则版本", "月搜索量", "ABA周排名", "PPC竞价", "来源ASIN", "数据质量告警", "备注"]
        values = [headers]
        for row in rows:
            data = _keyword_response(row, connection)
            values.append([data.get("keyword_raw"), data.get("keyword_normalized"), data.get("relevance_score"), f"{data.get('competitor_coverage', 0)}/{data.get('competitor_total', 0)}", data.get("match_strength"), data.get("suggested_action_label"), data.get("final_action_source"), data.get("rule_engine_version"), data.get("monthly_search_volume"), data.get("aba_weekly_rank"), data.get("ppc_bid_raw"), "/".join(data.get("related_asins") or []), ";".join(data.get("data_quality_flags") or []), data.get("notes")])
        if format == "csv":
            text_buffer = io.StringIO(newline="")
            writer = csv.writer(text_buffer)
            writer.writerows(values)
            content = text_buffer.getvalue().encode("utf-8-sig")
            media_type = "text/csv; charset=utf-8"
            filename = f"product-{product_id}-keywords.csv"
        else:
            try:
                from openpyxl import Workbook
            except ImportError:
                _api_error(503, "missing_dependency", "缺少 openpyxl，暂时只能导出 CSV")
            workbook = Workbook()
            sheet = workbook.active
            sheet.title = "关键词库"
            for row_values in values:
                sheet.append(row_values)
            workbook.save(output)
            content = output.getvalue()
            media_type = "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
            filename = f"product-{product_id}-keywords.xlsx"
    return StreamingResponse(io.BytesIO(content), media_type=media_type, headers={"Content-Disposition": f'attachment; filename="{filename}"'})


@app.get("/api/field-mapping")
def field_mapping() -> dict[str, Any]:
    from .importer import FIELD_ALIASES
    return {"fields": [{"field": field, "aliases": list(aliases)} for field, aliases in FIELD_ALIASES.items()], "required": ["keyword_raw"]}


# Import-time initialization makes ``uvicorn backend.app.main:app`` and simple
# scripts immediately usable; startup repeats it to cover process restarts.
init_db()
